import pptx
prs = pptx.Presentation(r'C:\Users\hp\Downloads\sih travel\SIH2026-IDEA-Presentation-Format.pptx')
for idx, sld in enumerate(prs.slides):
    if idx >= 6: break
    for s in sld.shapes:
        if 'Footer' in s.name or 'Number' in s.name or 'Oval' in s.name:
            t = s.text if s.has_text_frame else ''
            print(f'Slide {idx+1} {s.name}: text="{t}"')
